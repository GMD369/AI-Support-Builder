import io
from fastapi import HTTPException


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


def extract_text(filename: str, content: bytes) -> str:
    ext = _get_extension(filename)

    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
        )

    if ext in (".txt", ".md"):
        return _parse_text(content)
    if ext == ".pdf":
        return _parse_pdf(content)
    if ext == ".docx":
        return _parse_docx(content)
    if ext == ".xlsx":
        return _parse_xlsx(content)


def _get_extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _parse_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="Text file must be UTF-8 encoded")


def _parse_pdf(content: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(pages).strip()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse PDF: {e}")

    if not text:
        raise HTTPException(
            status_code=400,
            detail="PDF contains no extractable text (scanned/image-only PDFs are not supported)",
        )
    return text


def _parse_docx(content: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n\n".join(paragraphs).strip()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse DOCX: {e}")

    if not text:
        raise HTTPException(status_code=400, detail="DOCX file contains no readable text")
    return text


def _parse_xlsx(content: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join(str(c) for c in row if c is not None)
                if line.strip():
                    rows.append(line)
        text = "\n".join(rows).strip()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse XLSX: {e}")

    if not text:
        raise HTTPException(status_code=400, detail="XLSX file contains no readable data")
    return text
